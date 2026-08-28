from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "Event_Driven_Candidate_Sortie_RL_TITS_Draft.md"
OUTPUT = ROOT / "Event_Driven_Candidate_Sortie_RL_TITS_Draft.docx"

INK = "1B1F23"
MUTED = "5B6573"
BLUE = "1F4E79"
LIGHT_BLUE = "E8EEF5"
CAUTION = "FFF4CE"
RED = "9C2F2F"


def set_run_font(run, name="Times New Roman", size=10, bold=None, italic=None, color=INK):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "SimSun")
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=60, start=80, bottom=60, end=80):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_widths(table, widths_in):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    total = int(round(sum(widths_in) * 1440))
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "0")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_in:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(int(round(width * 1440))))
        grid.append(col)
    for row in table.rows:
        tr_pr = row._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)
        for idx, cell in enumerate(row.cells):
            width_dxa = int(round(widths_in[idx] * 1440))
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width_dxa))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(widths_in[idx])
            set_cell_margins(cell)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_columns(section, num=1, space_twips=360):
    sect_pr = section._sectPr
    cols = sect_pr.find(qn("w:cols"))
    if cols is None:
        cols = OxmlElement("w:cols")
        sect_pr.append(cols)
    cols.set(qn("w:num"), str(num))
    if num > 1:
        cols.set(qn("w:space"), str(space_twips))
        cols.set(qn("w:equalWidth"), "1")
    else:
        for key in ("space", "equalWidth"):
            cols.attrib.pop(qn(f"w:{key}"), None)


def setup_section(section, columns=1):
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.625)
    section.right_margin = Inches(0.625)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.3)
    set_columns(section, columns)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr, fld_char2])
    set_run_font(run, size=8, color=MUTED)


def add_furniture(section):
    header = section.header
    p = header.paragraphs[0]
    p.text = "AUTHOR REVIEW DRAFT · Target: IEEE Transactions on Intelligent Transportation Systems"
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    for run in p.runs:
        set_run_font(run, name="Arial", size=7.5, bold=True, color=MUTED)
    footer = section.footer
    add_page_number(footer.paragraphs[0])


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "SimSun")
    normal.font.size = Pt(10)
    pf = normal.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.space_before = Pt(0)
    pf.space_after = Pt(3)
    pf.line_spacing = 1.0

    h1 = doc.styles["Heading 1"]
    h1.font.name = "Times New Roman"
    h1._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    h1._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    h1.font.size = Pt(10)
    h1.font.bold = True
    h1.font.color.rgb = RGBColor.from_string(INK)
    h1.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    h1.paragraph_format.space_before = Pt(10)
    h1.paragraph_format.space_after = Pt(4)
    h1.paragraph_format.keep_with_next = True

    h2 = doc.styles["Heading 2"]
    h2.font.name = "Times New Roman"
    h2._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    h2._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    h2.font.size = Pt(10)
    h2.font.bold = False
    h2.font.italic = True
    h2.font.color.rgb = RGBColor.from_string(INK)
    h2.paragraph_format.space_before = Pt(7)
    h2.paragraph_format.space_after = Pt(3)
    h2.paragraph_format.keep_with_next = True

    for style_name in ("List Number", "List Bullet"):
        style = doc.styles[style_name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(10)
        style.paragraph_format.left_indent = Inches(0.30)
        style.paragraph_format.first_line_indent = Inches(-0.18)
        style.paragraph_format.space_after = Pt(3)
        style.paragraph_format.line_spacing = 1.0

    ref = doc.styles.add_style("IEEE Reference", 1)
    ref.font.name = "Times New Roman"
    ref._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    ref._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    ref.font.size = Pt(8)
    ref.paragraph_format.left_indent = Inches(0.2)
    ref.paragraph_format.first_line_indent = Inches(-0.2)
    ref.paragraph_format.space_after = Pt(1.5)
    ref.paragraph_format.line_spacing = 1.0


def clean_markdown(text):
    text = text.replace("**", "").replace("`", "")
    text = re.sub(r"(?<!\*)\*(?!\*)", "", text)
    text = text.replace("\\!", "")
    return text.strip()


def clean_inline_math(text):
    substitutions = {
        r"\bar{\Delta}": "Δ̄", r"\bar\theta": "θ̄", r"\mathbb{E}": "E",
        r"\arg\max": "arg max", r"\{": "{", r"\}": "}", r"\ldots": "…",
        r"\Delta": "Δ", r"\Gamma": "Γ", r"\theta": "θ", r"\gamma": "γ",
        r"\alpha": "α", r"\sigma": "σ", r"\rho": "ρ", r"\ell": "ℓ",
        r"\pi": "π", r"\in": "∈", r"\max": "max", r"\sum": "Σ",
        r"\left": "", r"\right": "", r"\!": "", r"\;": " ",
    }
    text = text.replace("$", "")
    for src, dst in substitutions.items():
        text = text.replace(src, dst)
    sub_map = str.maketrans("0123456789+-=()aikd", "₀₁₂₃₄₅₆₇₈₉₊₋₌₍₎ₐᵢₖᵈ")
    sup_map = str.maketrans("0123456789+-=()", "⁰¹²³⁴⁵⁶⁷⁸⁹⁺⁻⁼⁽⁾")

    def sub_braced(match):
        value = match.group(1)
        translated = value.translate(sub_map)
        return translated if translated != value else "_" + value

    def sup_braced(match):
        value = match.group(1)
        translated = value.translate(sup_map)
        return translated if translated != value else "^" + value

    text = re.sub(r"_\{([^{}]+)\}", sub_braced, text)
    text = re.sub(r"\^\{([^{}]+)\}", sup_braced, text)
    text = re.sub(r"_([0123456789aikd])", lambda m: m.group(1).translate(sub_map), text)
    text = re.sub(r"\^([0123456789+-])", lambda m: m.group(1).translate(sup_map), text)
    text = text.replace("\\", "")
    return text


def add_text_paragraph(doc, text, style=None, size=10, align=None, italic=False, bold=False, color=INK):
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    p.paragraph_format.widow_control = True
    run = p.add_run(clean_inline_math(clean_markdown(text)))
    set_run_font(run, size=size, italic=italic, bold=bold, color=color)
    return p


def add_equation(doc, latex):
    mapping = {
        r"\tau_i=(\ell_i,r_i,d_i,p_i,s_i,c_i),": "τᵢ = (ℓᵢ, rᵢ, dᵢ, pᵢ, sᵢ, cᵢ)",
        r"a_k=(g_k,\sigma_k,\rho_k),": "aₖ = (gₖ, σₖ, ρₖ)",
        r"\max_\pi \; \mathbb{E}_\pi\!\left[\sum_{k=0}^{K-1}\Gamma_k R_k\right],": "maxπ  Eπ [ Σₖ₌₀ᴷ⁻¹ Γₖ Rₖ ]",
        r"\gamma_k=\gamma^{\Delta t_k/\bar{\Delta}}.": "γₖ = γ^(Δtₖ / Δ̄)",
    }
    compact = latex.strip()
    if compact.startswith("a_k=\\arg"):
        value = "aₖ = arg max₍a : Mₖ,ₐ = 1₎ Qθ(sₖ, a)"
    elif compact.startswith("y_k=R_k"):
        value = "yₖ = Rₖ + (1 − zₖ) γₖ Qθ̄(sₖ₊₁, arg max₍a : Mₖ₊₁,ₐ = 1₎ Qθ(sₖ₊₁, a))"
    else:
        value = mapping.get(compact, compact)
    value = clean_inline_math(value).replace("\\!", "")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(value)
    set_run_font(run, name="Cambria Math", size=9.5, italic=True)


def add_table(doc, rows):
    col_count = len(rows[0])
    table = doc.add_table(rows=1, cols=col_count)
    table.style = "Table Grid"
    if col_count == 2:
        widths = [2.0, 5.1]
    elif col_count == 4:
        widths = [2.0, 1.35, 1.35, 2.4]
    elif col_count == 5:
        widths = [2.1, 1.25, 1.25, 1.25, 1.25]
    else:
        widths = [7.1 / col_count] * col_count

    header = table.rows[0]
    set_repeat_table_header(header)
    for idx, value in enumerate(rows[0]):
        cell = header.cells[idx]
        cell.text = ""
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        shade_cell(cell, BLUE)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(clean_inline_math(clean_markdown(value)))
        set_run_font(r, name="Arial", size=7.2, bold=True, color="FFFFFF")

    for ridx, values in enumerate(rows[1:], start=1):
        row = table.add_row()
        if ridx % 2 == 0:
            for cell in row.cells:
                shade_cell(cell, "F7F9FB")
        for idx, value in enumerate(values):
            cell = row.cells[idx]
            cell.text = ""
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if idx == 0 else WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(clean_inline_math(clean_markdown(value)))
            set_run_font(r, size=7.2)
    set_table_widths(table, widths)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    return table


def add_figure(doc, path_text):
    path = (ROOT / path_text).resolve()
    if not path.exists():
        raise FileNotFoundError(path)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    run = p.add_run()
    run.add_picture(str(path), width=Inches(6.85))


def add_review_callout(doc):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    cell = table.cell(0, 0)
    shade_cell(cell, CAUTION)
    set_cell_margins(cell, top=80, bottom=80, start=120, end=120)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("AUTHOR REVIEW DRAFT — NOT READY FOR SUBMISSION\n")
    set_run_font(r, name="Arial", size=9, bold=True, color=RED)
    r = p.add_run("Evidence is preliminary: 3-seed 20k screening + 1-seed 100k discount study. Missing final baselines and long-run multi-seed validation are marked in the text.")
    set_run_font(r, name="Arial", size=8, color=INK)
    set_table_widths(table, [7.0])


def main():
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    doc = Document()
    doc.core_properties.title = "Event-Driven Candidate-Sortie Reinforcement Learning for Dynamic CAV-UAV Cooperative Traffic Patrol"
    doc.core_properties.subject = "Author review draft for IEEE Transactions on Intelligent Transportation Systems"
    doc.core_properties.author = "Yuxiang Ouyang; Shijia Zhang; Yikang Rui; Ran Bin"
    configure_styles(doc)
    setup_section(doc.sections[0], columns=1)
    add_furniture(doc.sections[0])
    add_review_callout(doc)

    i = 0
    in_comment = False
    in_abstract = False
    in_references = False
    wide_mode = False
    table_pending_note = False
    current_columns = 1

    def new_section(columns, start=WD_SECTION.CONTINUOUS):
        nonlocal current_columns
        if columns == current_columns:
            return
        section = doc.add_section(start)
        setup_section(section, columns=columns)
        section.header.is_linked_to_previous = True
        section.footer.is_linked_to_previous = True
        current_columns = columns

    while i < len(lines):
        raw = lines[i]
        line = raw.strip()
        if line.startswith("<!--"):
            in_comment = True
        if in_comment:
            if line.endswith("-->"):
                in_comment = False
            i += 1
            continue
        if not line:
            i += 1
            continue

        if table_pending_note and not line.startswith("*"):
            new_section(2)
            table_pending_note = False
            wide_mode = False

        if line.startswith("# "):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(6)
            r = p.add_run(line[2:])
            set_run_font(r, size=20, bold=True)
            i += 1
            continue

        if line == "## Abstract":
            in_abstract = True
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(5)
            p.paragraph_format.space_after = Pt(3)
            r = p.add_run("Abstract")
            set_run_font(r, size=9, bold=True, italic=True)
            i += 1
            continue

        if line == "## I. Introduction":
            in_abstract = False
            new_section(2)
            p = doc.add_paragraph("I. INTRODUCTION", style="Heading 1")
            i += 1
            continue

        if line.startswith("## Author-Only Chinese Abstract"):
            new_section(1, WD_SECTION.NEW_PAGE)
            p = doc.add_paragraph("AUTHOR-ONLY CHINESE ABSTRACT (NOT FOR SUBMISSION)", style="Heading 1")
            i += 1
            continue

        if line.startswith("## Author Checklist"):
            p = doc.add_paragraph("AUTHOR CHECKLIST BEFORE T-ITS SUBMISSION", style="Heading 1")
            i += 1
            continue

        if line.startswith("## "):
            heading = clean_markdown(line[3:])
            if heading == "References":
                in_references = True
            p = doc.add_paragraph(heading.upper(), style="Heading 1")
            i += 1
            continue

        if line.startswith("### "):
            p = doc.add_paragraph(clean_markdown(line[4:]), style="Heading 2")
            i += 1
            continue

        if line.startswith("**TABLE"):
            new_section(1)
            wide_mode = True
            p = add_text_paragraph(doc, line, size=8.2, align=WD_ALIGN_PARAGRAPH.CENTER, bold=True)
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(2)
            i += 1
            continue

        if line.startswith("!["):
            match = re.match(r"!\[[^\]]*\]\(([^)]+)\)", line)
            if match:
                new_section(1)
                wide_mode = True
                add_figure(doc, match.group(1))
            i += 1
            continue

        if line.startswith("**Fig."):
            p = add_text_paragraph(doc, line, size=8, align=WD_ALIGN_PARAGRAPH.CENTER)
            p.paragraph_format.keep_with_next = False
            p.paragraph_format.space_after = Pt(5)
            if wide_mode:
                new_section(2)
                wide_mode = False
            i += 1
            continue

        if line.startswith("|"):
            table_rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [cell.strip() for cell in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{3,}:?", c) for c in cells):
                    table_rows.append(cells)
                i += 1
            add_table(doc, table_rows)
            table_pending_note = True
            continue

        if line == "$$":
            equation = []
            i += 1
            while i < len(lines) and lines[i].strip() != "$$":
                equation.append(lines[i].strip())
                i += 1
            add_equation(doc, " ".join(equation))
            i += 1
            continue

        if re.match(r"^\d+\. ", line):
            p = doc.add_paragraph(style="List Number")
            r = p.add_run(re.sub(r"^\d+\. ", "", line))
            set_run_font(r, size=10)
            i += 1
            continue

        if line.startswith("- [ ]"):
            p = doc.add_paragraph(style="List Bullet")
            r = p.add_run("☐ " + line[5:].strip())
            set_run_font(r, size=10)
            i += 1
            continue

        if line == "---":
            i += 1
            continue

        if table_pending_note and line.startswith("*"):
            p = add_text_paragraph(doc, line, size=7.5, italic=True, color=MUTED)
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(4)
            new_section(2)
            table_pending_note = False
            wide_mode = False
            i += 1
            continue

        if in_references and re.match(r"^\[\d+\]", line):
            p = add_text_paragraph(doc, line, style="IEEE Reference", size=8)
            i += 1
            continue

        if line.startswith("**Index Terms"):
            p = add_text_paragraph(doc, line, size=9)
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(5)
            i += 1
            continue

        if i > 0 and lines[i - 1].strip().startswith("# "):
            p = add_text_paragraph(doc, line, size=10.5, align=WD_ALIGN_PARAGRAPH.CENTER)
            p.paragraph_format.space_after = Pt(4)
            i += 1
            continue

        if in_abstract:
            p = add_text_paragraph(doc, line, size=9)
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.right_indent = Inches(0.25)
            p.paragraph_format.space_after = Pt(3)
            i += 1
            continue

        p = add_text_paragraph(doc, line, size=9.5 if current_columns == 2 else 10)
        if "[AUTHOR TO" in line or "PROVISIONAL" in line:
            for run in p.runs:
                run.font.color.rgb = RGBColor.from_string(RED)
        i += 1

    for section in doc.sections:
        setup_section(section, columns=int(section._sectPr.find(qn("w:cols")).get(qn("w:num"), "1")))
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
